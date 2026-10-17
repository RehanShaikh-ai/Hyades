"""DOCX file extractor.

Canonical module per CONTRACT v0.4.2 §11.2.
Extracts readable text from paragraphs and tables while enforcing archive limits.
"""

import io
from typing import Any

from app.services.extractors.base import BaseExtractor, ExtractedContent, safe_open_zip


class DocxExtractor(BaseExtractor):
    """Extractor for DOCX documents."""

    def extract(self, file_bytes: bytes, filename: str = "", **kwargs: Any) -> ExtractedContent:
        # Enforce zip-bomb archive limits before parsing with docx
        safe_open_zip(file_bytes)

        try:
            import docx
        except ImportError as err:
            raise ValueError("python-docx is not installed in the runtime environment.") from err

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {e}") from e

        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract table text
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    tables_text.append(" | ".join(row_cells))

        all_blocks = paragraphs_text + tables_text
        full_text = "\n\n".join(all_blocks)

        word_count = len(full_text.split()) if full_text else 0
        metadata: dict[str, Any] = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "word_count": word_count,
            "format": "docx",
        }

        # Core properties if present
        if hasattr(doc, "core_properties"):
            props = doc.core_properties
            if props.title:
                metadata["title"] = props.title
            if props.author:
                metadata["author"] = props.author

        return ExtractedContent(text=full_text, metadata=metadata)
