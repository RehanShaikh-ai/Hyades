import io
import uuid
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.user import User
from app.models.workspace import Workspace
from app.models.source import Source
from app.services.document_preview_service import (
    generate_source_preview,
    sanitize_markdown,
    PreviewConversionError,
)



@pytest.fixture
def test_user_and_workspace(db_session: Session) -> tuple[User, Workspace]:
    user = User(id=uuid.uuid4(), display_name="Test Importer")
    db_session.add(user)
    db_session.flush()
    ws = Workspace(id=uuid.uuid4(), name="Test Vault", owner_id=user.id)
    db_session.add(ws)
    db_session.commit()
    return user, ws


def test_sanitize_markdown_removes_scripts_and_event_handlers():
    """Verify HTML script tags, event handlers, and javascript URIs are strictly stripped."""
    dangerous = """# Safe Heading
<script>alert('pwned');</script>
<p onclick="stealCookies()">Paragraph with event</p>
<iframe src="evil.com"></iframe>
[Link](javascript:alert(1))
"""
    clean = sanitize_markdown(dangerous)
    assert "<script>" not in clean
    assert "alert('pwned')" not in clean
    assert "<iframe" not in clean
    assert "onclick=" not in clean
    assert "javascript:" not in clean
    assert "Safe Heading" in clean
    assert "Paragraph with event" in clean


def test_generate_source_preview_markdown_file():
    """Verify native markdown document returns correctly formatted markdown."""
    sid = uuid.uuid4()
    md_content = b"""# Celestial Observations
The Hyades cluster is the nearest open cluster to the Sun.
* Star 1: Aldebaran alignment
* Star 2: Theta Tauri
"""
    res = generate_source_preview(
        source_id=sid,
        original_path="astronomy/hyades.md",
        source_type="markdown",
        file_bytes=md_content,
        raw_metadata={"title": "Celestial Observations"},
    )
    assert res["source_id"] == str(sid)
    assert res["title"] == "Celestial Observations"
    assert res["format"] == "md"
    assert "The Hyades cluster is the nearest" in res["markdown_content"]
    assert res["converter"] == "native_markdown"


def test_generate_source_preview_docx_file():
    """Verify supported non-markdown format (DOCX) is converted into structured markdown."""
    import docx
    doc = docx.Document()
    doc.add_heading("Research Notes on Knowledge Graphs", level=1)
    doc.add_paragraph("Knowledge graphs represent structured relationships between entities.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    sid = uuid.uuid4()
    res = generate_source_preview(
        source_id=sid,
        original_path="research_notes.docx",
        source_type="docx",
        file_bytes=docx_bytes,
        raw_metadata={"title": "Research Notes on Knowledge Graphs"},
    )
    assert res["source_id"] == str(sid)
    assert res["format"] == "docx"
    assert res["is_converted"] is True
    assert "Research Notes on Knowledge Graphs" in res["markdown_content"]
    assert "Knowledge graphs represent structured relationships" in res["markdown_content"]


def test_generate_source_preview_unsupported_format_raises_error():
    """Verify unsupported binary/executable file raises clear PreviewConversionError."""
    sid = uuid.uuid4()
    binary_bytes = b"\x7fELF\x02\x01\x01\x00"
    with pytest.raises(PreviewConversionError) as exc_info:
        generate_source_preview(
            source_id=sid,
            original_path="malicious_binary.exe",
            source_type="binary",
            file_bytes=binary_bytes,
        )
    assert "not supported for text preview" in str(exc_info.value)


def test_source_document_preview_api_endpoint(
    client: TestClient, test_user_and_workspace: tuple[User, Workspace]
):
    """Test full GET /api/v1/sources/{source_id}/preview endpoint with markdown upload."""
    _, ws = test_user_and_workspace
    content = b"""# Epistemological Notes
Dialectical synthesis between personal observations and external knowledge.
"""
    files = {"file": ("epistemology.md", io.BytesIO(content), "text/markdown")}
    upload_res = client.post(f"/api/v1/workspaces/{ws.id}/sources/upload", files=files)
    assert upload_res.status_code == 201
    source_id = upload_res.json()["id"]

    # Preview endpoint call
    prev_res = client.get(f"/api/v1/sources/{source_id}/preview")
    assert prev_res.status_code == 200
    data = prev_res.json()
    assert data["source_id"] == source_id
    assert "Epistemological Notes" in data["markdown_content"]
    assert data["format"] == "md"
